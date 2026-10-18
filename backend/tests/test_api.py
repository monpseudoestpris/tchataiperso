import io
import json
from contextlib import contextmanager

import pytest
from fastapi.testclient import TestClient
from docx import Document
from PIL import Image
from pydantic import SecretStr
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app import providers
from app.auth import hash_password
from app.config import Settings
from app.main import create_app
from app.models import Conversation, Message, User


@pytest.fixture
def settings(tmp_path):
    return Settings(_env_file=None, app_data_dir=tmp_path, demo_mode=True)


@contextmanager
def authenticated_client(settings, app=None):
    app = app or create_app(settings)
    origin = settings.origins[0]
    from app.models import Base
    Base.metadata.create_all(app.state.engine)
    with app.state.sessions() as db:
        db.add(User(username="tester", password_hash=hash_password("correct horse battery"), is_admin=True))
        db.commit()
    with TestClient(app, base_url=origin, headers={"Origin": origin}) as value:
        response = value.post("/api/auth/login", json={"username": "tester", "password": "correct horse battery"})
        assert response.status_code == 200
        yield value


@pytest.fixture
def client(settings):
    with authenticated_client(settings) as value:
        yield value


def events(response):
    result = []
    for block in response.text.strip().split("\n\n"):
        lines = block.splitlines()
        name = next(line[7:] for line in lines if line.startswith("event: "))
        data = json.loads(next(line[6:] for line in lines if line.startswith("data: ")))
        result.append((name, data))
    return result


def conversation(client, **body):
    response = client.post("/api/conversations", json=body)
    assert response.status_code == 201
    return response.json()["id"]


def test_config_and_disabled_credentials(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    config = client.get("/api/config").json()
    assert config["demo_mode"] is True
    assert config["features"] == {"chat": True, "web_search": False, "ocr": False, "image_generation": False}
    assert "api_key" not in json.dumps(config).lower()


def test_app_fails_with_actionable_error_without_admin(tmp_path):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, demo_mode=True)
    with pytest.raises(RuntimeError, match="Aucun administrateur actif"):
        with TestClient(create_app(settings)):
            pass


def test_wildcard_cors_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="wildcard est interdit"):
        Settings(_env_file=None, app_data_dir=tmp_path, cors_origins="*")


def test_private_api_requires_authenticated_session(client):
    origin = client.app.state.settings.origins[0]
    with TestClient(client.app, base_url=origin, headers={"Origin": origin}) as anonymous:
        assert anonymous.get("/api/conversations").status_code == 401
        assert anonymous.get("/api/memories").status_code == 401
        assert anonymous.post("/api/conversations", json={}).status_code == 401
        assert anonymous.post("/api/admin/invitations").status_code == 401


def test_session_cookie_is_http_only_and_same_site(client):
    response = client.post("/api/auth/login", json={
        "username": "tester", "password": "correct horse battery",
    })
    cookie = response.headers["set-cookie"].lower()
    assert "httponly" in cookie
    assert "samesite=lax" in cookie
    assert "path=/" in cookie


def test_logout_revokes_session_and_login_is_rate_limited(client):
    assert client.get("/api/auth/session").status_code == 200
    assert client.post("/api/auth/logout").status_code == 204
    assert client.get("/api/auth/session").status_code == 401

    origin = client.app.state.settings.origins[0]
    with TestClient(client.app, base_url=origin, headers={"Origin": origin}) as anonymous:
        for _ in range(10):
            assert anonymous.post("/api/auth/login", json={
                "username": "tester", "password": "incorrect password here",
            }).status_code == 401
        assert anonymous.post("/api/auth/login", json={
            "username": "tester", "password": "incorrect password here",
        }).status_code == 429


def test_mutations_reject_untrusted_origins(client):
    response = client.post("/api/memories", json={"content": "refused"},
                           headers={"Origin": "https://untrusted.example"})
    assert response.status_code == 403
    with TestClient(client.app, base_url=client.app.state.settings.origins[0]) as no_origin:
        assert no_origin.post("/api/auth/login", json={
            "username": "tester", "password": "correct horse battery",
        }).status_code == 403


def test_invited_users_cannot_access_another_users_data(client):
    owned_conversation = conversation(client)
    owned_memory = client.post("/api/memories", json={"content": "Privé pour tester"}).json()
    owned_attachment = client.post("/api/attachments", files={"file": ("secret.txt", b"secret", "text/plain")}).json()
    invitation = client.post("/api/admin/invitations")
    assert invitation.status_code == 201

    origin = client.app.state.settings.origins[0]
    with TestClient(client.app, base_url=origin, headers={"Origin": origin}) as other:
        registration = other.post("/api/auth/register", json={
            "token": invitation.json()["token"], "username": "second-user", "password": "another secure password",
        })
        assert registration.status_code == 201
        assert other.post("/api/auth/register", json={
            "token": invitation.json()["token"], "username": "third-user", "password": "another secure password",
        }).status_code == 400
        assert other.post("/api/auth/login", json={
            "username": "second-user", "password": "another secure password",
        }).status_code == 200
        assert other.get("/api/conversations").json() == []
        assert other.get(f"/api/conversations/{owned_conversation}").status_code == 404
        assert other.patch(f"/api/conversations/{owned_conversation}", json={"title": "Volé"}).status_code == 404
        assert other.delete(f"/api/conversations/{owned_conversation}").status_code == 404
        assert other.get("/api/memories").json() == []
        assert other.delete(f"/api/memories/{owned_memory['id']}").status_code == 404
        assert other.get(owned_attachment["url"]).status_code == 404
        assert other.delete(f"/api/attachments/{owned_attachment['id']}").status_code == 404

        other_conversation = conversation(other)
        response = other.post(f"/api/conversations/{other_conversation}/messages", json={
            "content": "Joindre le fichier d'un autre compte", "attachment_ids": [owned_attachment["id"]],
        })
        assert response.status_code == 422
        assert client.get(owned_attachment["url"]).status_code == 200
        assert other.post("/api/admin/invitations").status_code == 403
        user_id = registration.json()["id"]
        assert client.patch(f"/api/admin/users/{user_id}", json={"active": False}).status_code == 200
        assert other.get("/api/memories").status_code == 401
        assert client.patch(f"/api/admin/users/{user_id}", json={"active": True}).status_code == 200
        assert other.post("/api/auth/login", json={
            "username": "second-user", "password": "another secure password",
        }).status_code == 200


def test_demo_stream_history_and_crud(client):
    item = conversation(client)
    sent = client.post(f"/api/conversations/{item}/messages", json={"content": "Bonjour"})
    stream = events(sent)
    assert stream[0][0] == "meta"
    assert stream[-1][0] == "done"
    assert "Mode démonstration" in "".join(data["text"] for name, data in stream if name == "delta")
    detail = client.get(f"/api/conversations/{item}").json()
    assert [message["role"] for message in detail["messages"]] == ["user", "assistant"]
    assert detail["title"] == "Bonjour"
    assert client.patch(f"/api/conversations/{item}", json={"title": "Test"}).json()["title"] == "Test"
    assert client.delete(f"/api/conversations/{item}").status_code == 204
    assert client.get(f"/api/conversations/{item}").status_code == 404


def test_model_and_feature_validation(client):
    assert client.post("/api/conversations", json={"model": "unknown"}).status_code == 422
    item = conversation(client)
    route = f"/api/conversations/{item}/messages"
    assert client.post(route, json={"content": "test", "model": "unknown"}).status_code == 422
    assert client.post(route, json={"content": ""}).status_code == 422
    assert client.post(route, json={"content": "test", "web_search": True}).status_code == 503
    assert client.post(f"/api/conversations/{item}/images", json={"prompt": "a cat"}).status_code == 503


def test_upload_attachment_extraction_history_and_deletion(client, settings):
    upload = client.post("/api/attachments", files={"file": ("../notes.md", b"# Document\nUne information", "text/plain")})
    assert upload.status_code == 201
    attachment = upload.json()
    assert attachment["filename"] == "notes.md"
    assert attachment["extracted_chars"] > 0
    file_response = client.get(attachment["url"])
    assert file_response.content.startswith(b"# Document")
    assert file_response.headers["cache-control"] == "no-store"
    item = conversation(client)
    response = client.post(f"/api/conversations/{item}/messages", json={"content": "Lis", "attachment_ids": [attachment["id"]]})
    assert events(response)[-1][0] == "done"
    detail = client.get(f"/api/conversations/{item}").json()
    assert detail["messages"][0]["attachments"][0]["id"] == attachment["id"]
    assert client.post(f"/api/conversations/{item}/messages", json={"content": "Bis", "attachment_ids": [attachment["id"]]}).status_code == 422
    assert client.delete(f"/api/conversations/{item}").status_code == 204
    assert client.get(attachment["url"]).status_code == 404
    assert list((settings.app_data_dir / "uploads").iterdir()) == []


def test_scanned_pdf_requires_ocr_and_unsafe_upload_rejected(client):
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    content = io.BytesIO()
    pdf.write(content)
    response = client.post("/api/attachments", files={"file": ("scan.pdf", content.getvalue(), "application/pdf")})
    assert response.status_code == 422
    assert "MISTRAL_API_KEY" in response.json()["detail"]
    assert client.post("/api/attachments", files={"file": ("code.html", b"<script>alert(1)</script>", "text/html")}).status_code == 415
    assert client.post("/api/attachments", files={"file": ("fake.png", b"bad image", "image/png")}).status_code == 422


def test_upload_limits(client, settings):
    settings.max_upload_mb = 1
    response = client.post("/api/attachments", files={"file": ("large.txt", b"x" * (1024 * 1024 + 1), "text/plain")})
    assert response.status_code == 413


def test_explicit_memory_is_persisted_and_deleted(client):
    item = client.post("/api/memories", json={"content": "Je préfère le français"}).json()
    assert client.get("/api/memories").json() == [item]
    assert client.delete(f"/api/memories/{item['id']}").status_code == 204
    assert client.get("/api/memories").json() == []
    assert client.post("/api/memories", json={"content": "  "}).status_code == 422


def test_no_key_never_silently_uses_demo(tmp_path):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, demo_mode=False)
    with authenticated_client(settings) as client:
        assert client.get("/api/config").json()["features"]["chat"] is False
        item = conversation(client)
        response = client.post(f"/api/conversations/{item}/messages", json={"content": "hello"})
        assert response.status_code == 503
        assert client.get(f"/api/conversations/{item}").json()["messages"] == []


def test_mock_real_chat_injects_memory_docs_and_sources(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, demo_mode=False, llm_api_key="secret", brave_api_key="brave")
    captured = {}

    async def search(_settings, query):
        assert query == "Compare les infos"
        return [{"title": "Référence", "url": "https://example.org", "snippet": "Texte web"}]

    async def stream(_settings, messages, model):
        captured["messages"] = messages
        yield "Réponse "
        yield "réelle simulée"

    monkeypatch.setattr(providers, "search_web", search)
    monkeypatch.setattr(providers, "stream_chat", stream)
    with authenticated_client(settings) as client:
        client.post("/api/memories", json={"content": "Répondre en français"})
        attachment = client.post("/api/attachments", files={"file": ("facts.txt", b"Information document", "text/plain")}).json()
        item = conversation(client)
        response = client.post(f"/api/conversations/{item}/messages", json={"content": "Compare les infos", "web_search": True, "attachment_ids": [attachment["id"]]})
        assert events(response)[-1][0] == "done"
        payload = json.dumps(captured["messages"], ensure_ascii=False)
        assert "Répondre en français" in payload
        assert "Information document" in payload
        assert "https://example.org" in payload
        assert client.get(f"/api/conversations/{item}").json()["messages"][-1]["sources"][0]["title"] == "Référence"


def test_stream_failure_is_explicit_unlocks_and_preserves_user(client, settings, monkeypatch):
    settings.demo_mode = False
    settings.llm_api_key = SecretStr("secret")

    async def failing(*args):
        yield "Début"
        raise providers.ProviderError("Quota épuisé")

    monkeypatch.setattr(providers, "stream_chat", failing)
    item = conversation(client)
    response = client.post(f"/api/conversations/{item}/messages", json={"content": "hello"})
    assert events(response)[-1] == ("error", {"message": "Quota épuisé"})
    assert [value["role"] for value in client.get(f"/api/conversations/{item}").json()["messages"]] == ["user"]
    assert client.patch(f"/api/conversations/{item}", json={"title": "unlocked"}).status_code == 200


def test_summary_retains_complete_history_and_excludes_global_memory(client, settings):
    settings.summary_every_messages = 8
    client.post("/api/memories", json={"content": "GLOBAL_PRIVATE_MEMORY"})
    item = conversation(client)
    for index in range(5):
        response = client.post(f"/api/conversations/{item}/messages", json={"content": f"Échange numéro {index}"})
        assert events(response)[-1][0] == "done"
    detail = client.get(f"/api/conversations/{item}").json()
    assert len(detail["messages"]) == 10
    assert detail["summary"]
    assert "GLOBAL_PRIVATE_MEMORY" not in detail["summary"]


def test_image_model_and_ocr_fallback(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, demo_mode=False, llm_api_key="secret",
                        mistral_api_key="ocr", llm_models='[{"id":"text","name":"Text","vision":false}]', default_model="text")
    content = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(content, format="PNG")
    captured = {}

    async def ocr(_settings, data, mime):
        assert mime == "image/png"
        return "TEXTE OCR"

    async def stream(_settings, messages, model):
        captured["messages"] = messages
        yield "Texte lu"

    monkeypatch.setattr(providers, "extract_ocr", ocr)
    monkeypatch.setattr(providers, "stream_chat", stream)
    with authenticated_client(settings) as client:
        upload = client.post("/api/attachments", files={"file": ("image.png", content.getvalue(), "image/png")}).json()
        item = conversation(client)
        response = client.post(f"/api/conversations/{item}/messages", json={"content": "lis", "attachment_ids": [upload["id"]]})
        assert events(response)[-1][0] == "done"
        assert "TEXTE OCR" in json.dumps(captured)


def test_mock_image_generation_saved_in_history(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, ideogram_api_key="image", demo_mode=False)

    async def generate(_settings, prompt, ratio):
        assert prompt == "Un chat" and ratio == "1:1"
        return b"image-data", "image/png"

    monkeypatch.setattr(providers, "generate_image", generate)
    with authenticated_client(settings) as client:
        item = conversation(client)
        response = client.post(f"/api/conversations/{item}/images", json={"prompt": "Un chat"})
        assert response.status_code == 201
        image = response.json()
        assert image["attachments"][0]["mime_type"] == "image/png"
        assert client.get(image["attachments"][0]["url"]).content == b"image-data"
        assert len(client.get(f"/api/conversations/{item}").json()["messages"]) == 2


def test_summary_includes_old_document_text_and_preserves_unsummarized_turns(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, app_data_dir=tmp_path, llm_api_key="key", summary_every_messages=8)
    captured = {}

    async def summarize(_settings, model, previous, transcript):
        captured["transcript"] = transcript
        return "Résumé fiable"

    async def stream(_settings, messages, model):
        captured["context"] = json.dumps(messages, ensure_ascii=False)
        yield "Réponse"

    monkeypatch.setattr(providers, "summarize", summarize)
    monkeypatch.setattr(providers, "stream_chat", stream)
    app = create_app(settings)
    with authenticated_client(settings, app) as client:
        item = conversation(client)
        document = client.post("/api/attachments", files={"file": ("facts.txt", b"OLD_DOCUMENT_FACT", "text/plain")}).json()
        client.post(f"/api/conversations/{item}/messages", json={"content": "Premier échange", "attachment_ids": [document["id"]]})
        with app.state.sessions() as db:
            # Long previous turns force a bounded partial summary batch.
            for index in range(14):
                db.add(Message(conversation_id=item, role="user" if index % 2 == 0 else "assistant",
                               content=f"MARKER_{index} " + "x" * 6000))
            db.commit()
        response = client.post(f"/api/conversations/{item}/messages", json={"content": "Nouvelle question"})
        assert events(response)[-1][0] == "done"
        assert "OLD_DOCUMENT_FACT" in captured["transcript"]
        with app.state.sessions() as db:
            stored = db.get(Conversation, item)
            assert stored.summary_through < len(stored.messages) - 7
            first_unsummarized = stored.messages[stored.summary_through].content.split(" ")[0]
        assert first_unsummarized not in captured["transcript"]
        assert first_unsummarized in captured["context"]


def test_real_pdf_and_docx_extraction(client):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 200 Td (PDF_DOCUMENT_FACT) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf = io.BytesIO()
    writer.write(pdf)
    pdf_response = client.post("/api/attachments", files={"file": ("facts.pdf", pdf.getvalue(), "application/pdf")})
    assert pdf_response.status_code == 201
    assert pdf_response.json()["extracted_chars"] >= len("PDF_DOCUMENT_FACT")

    document = Document()
    document.add_paragraph("DOCX_DOCUMENT_FACT")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TABLE_DOCUMENT_FACT"
    docx = io.BytesIO()
    document.save(docx)
    docx_response = client.post("/api/attachments", files={"file": ("facts.docx", docx.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert docx_response.status_code == 201
    assert docx_response.json()["extracted_chars"] >= len("DOCX_DOCUMENT_FACT\nTABLE_DOCUMENT_FACT")

    item = conversation(client)
    client.post(f"/api/conversations/{item}/messages", json={"content": "Résume", "attachment_ids": [pdf_response.json()["id"], docx_response.json()["id"]]})
    with client.app.state.sessions() as db:
        stored = db.get(Conversation, item)
        texts = [attachment.extracted_text for attachment in stored.messages[0].attachments]
    assert any("PDF_DOCUMENT_FACT" in value for value in texts)
    assert any("DOCX_DOCUMENT_FACT" in value and "TABLE_DOCUMENT_FACT" in value for value in texts)


def test_orphan_upload_delete_and_linked_file_protected(client):
    item = client.post("/api/attachments", files={"file": ("unused.txt", b"Unused", "text/plain")}).json()
    assert client.delete(f"/api/attachments/{item['id']}").status_code == 204
    assert client.get(item["url"]).status_code == 404
    attached = client.post("/api/attachments", files={"file": ("used.txt", b"Used", "text/plain")}).json()
    conv = conversation(client)
    client.post(f"/api/conversations/{conv}/messages", json={"content": "Lis", "attachment_ids": [attached["id"]]})
    assert client.delete(f"/api/attachments/{attached['id']}").status_code == 409


def test_five_images_rejected_explicitly(client):
    image = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    attachments = [client.post("/api/attachments", files={"file": (f"{index}.png", image.getvalue(), "image/png")}).json()["id"] for index in range(5)]
    item = conversation(client)
    response = client.post(f"/api/conversations/{item}/messages", json={"content": "Compare", "attachment_ids": attachments})
    assert response.status_code == 422
    assert "4 images" in response.json()["detail"]
    assert client.get(f"/api/conversations/{item}").json()["messages"] == []
