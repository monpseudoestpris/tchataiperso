import asyncio
import io
import re
import warnings
import zipfile
from contextlib import closing
from pathlib import Path
from threading import Lock

from docx import Document
from fastapi import HTTPException
from PIL import Image
from pypdf import PdfReader
import pypdfium2 as pdfium

from . import providers
from .config import Settings
from .schemas import DocumentExtraction, ExtractedPage

MAX_EXTRACTED_CHARS = 40000
_PDF_RENDER_LOCK = Lock()


def inspect_document(content: bytes, filename: str) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(content)) as picture:
                    picture.verify()
                    if picture.format not in {"PNG", "JPEG", "WEBP"}:
                        raise ValueError()
                    mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}[picture.format]
                    return mime, ""
        except Exception as exc:
            raise HTTPException(422, "Image invalide ou dimensions excessives. Formats acceptés : PNG, JPEG, WebP.") from exc
    if suffix == ".pdf":
        try:
            if not content.startswith(b"%PDF-"):
                raise ValueError()
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise HTTPException(422, "Déverrouillez le PDF avant de le joindre.")
            if len(reader.pages) > 300:
                raise HTTPException(422, "Le PDF dépasse la limite de 300 pages.")
            chunks: list[str] = []
            length = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                chunks.append(text)
                length += len(text)
                if length >= MAX_EXTRACTED_CHARS:
                    break
            return "application/pdf", "\n".join(chunks)[:MAX_EXTRACTED_CHARS]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "PDF illisible ou non pris en charge.") from exc
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError()
            document = Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "\n".join(paragraphs)[:MAX_EXTRACTED_CHARS]
        except Exception as exc:
            raise HTTPException(422, "Document Word invalide ou trop volumineux après décompression.") from exc
    if suffix in {".txt", ".md", ".csv", ".json"}:
        try:
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError()
            return "text/plain", text[:MAX_EXTRACTED_CHARS]
        except (UnicodeDecodeError, ValueError) as exc:
            raise HTTPException(422, "Le document texte doit être encodé en UTF-8.") from exc
    raise HTTPException(415, "Formats acceptés : PDF, DOCX, TXT, MD, CSV, JSON, PNG, JPEG et WebP.")


def extraction_text(extraction: dict | DocumentExtraction, limit: int | None = None) -> str:
    result = DocumentExtraction.model_validate(extraction)
    sections = []
    page_contents = []
    for page in result.pages:
        descriptions = [f"Figure {image.id} : {image.annotation.description}\n{image.annotation.visible_text}"
                        for image in page.images if image.annotation]
        if page.vision_description:
            descriptions.append(f"Description visuelle : {page.vision_description}")
        header = f"[Page {page.index + 1}]\n"
        description = "\n\n".join(descriptions)
        page_contents.append((header, page.markdown, description))
        sections.append(header + "\n\n".join(value for value in (page.markdown, description) if value))
    full = "\n\n".join(sections)
    if limit is None or len(full) <= limit:
        return full
    notice = "[Extraits limites par page ; contenu integral disponible dans la piece jointe.]\n"
    allowance = max(0, (limit - len(notice)) // len(sections) - 2)
    marker = " [...]"

    def excerpt(value: str, size: int) -> str:
        if len(value) <= size:
            return value
        return (value[:max(0, size - len(marker))] + marker)[:size]

    excerpts = []
    for header, markdown, description in page_contents:
        content_budget = max(0, allowance - len(header) - 2)
        description_budget = min(len(description), content_budget // 3) if markdown else content_budget
        text_budget = max(0, content_budget - description_budget)
        excerpts.append(header + excerpt(markdown, text_budget) + "\n\n" + excerpt(description, description_budget))
    return (notice + "\n\n".join(excerpts))[:limit]


def render_pdf_page(content: bytes, index: int) -> bytes:
    try:
        with _PDF_RENDER_LOCK, pdfium.PdfDocument(content) as document:
            with closing(document[index]) as page:
                width, height = page.get_size()
                if min(width, height) <= 0 or max(width, height) / min(width, height) > 100:
                    raise ValueError
                bitmap = page.render(scale=min(2, 2048 / max(width, height)))
                try:
                    with bitmap.to_pil() as image:
                        output = io.BytesIO()
                        image.save(output, format="PNG")
                        return output.getvalue()
                finally:
                    bitmap.close()
    except Exception:
        raise HTTPException(422, "Impossible de rendre cette page PDF pour la vision.") from None


def normalize_vision_image(content: bytes) -> bytes:
    with Image.open(io.BytesIO(content)) as image:
        image.thumbnail((2048, 2048))
        with image.convert("RGB") as picture:
            output = io.BytesIO()
            picture.save(output, format="JPEG", quality=90)
            return output.getvalue()


async def extract_document(settings: Settings, content: bytes, filename: str, mode: str = "auto") -> tuple[str, str, dict | None]:
    mime, text = await asyncio.to_thread(inspect_document, content, filename)
    visual = mime == "application/pdf" or mime.startswith("image/")
    if mime.startswith("image/") and len(content) > 8 * 1024 * 1024:
        raise HTTPException(413, "Les images sont limitées à 8 Mo pour la vision")

    # Mode Vision explicite : analyse visuelle détaillée par le modèle de vision
    if mode in ("vision", "image"):
        if not visual:
            raise HTTPException(422, "L'analyse visuelle détaillée s'applique uniquement aux images (PNG, JPEG, WebP) et aux PDF.")
        if settings.demo_mode:
            description = f"Description visuelle détaillée simulée pour {filename}."
            extraction = DocumentExtraction(
                model="demo-vision",
                pages=[ExtractedPage(index=0, markdown=description, vision_description=description, method="vision")],
                warnings=[],
            )
            return mime, description, extraction.model_dump()
        if not settings.vision_model:
            raise HTTPException(422, "Configurez VISION_MODEL pour analyser les images avec le modèle de vision.")
        try:
            if mime == "application/pdf":
                image = await asyncio.to_thread(render_pdf_page, content, 0)
                image_mime = "image/png"
            else:
                image = await asyncio.to_thread(normalize_vision_image, content)
                image_mime = "image/jpeg"
            async with asyncio.timeout(300):
                description = await providers.describe_image(settings, image, image_mime)
            extraction = DocumentExtraction(
                model=settings.vision_model,
                pages=[ExtractedPage(index=0, markdown=description, vision_description=description, method="vision")],
                warnings=[],
            )
            return mime, description, extraction.model_dump()
        except TimeoutError:
            raise providers.ProviderError("L'analyse visuelle a dépassé le délai autorisé de 5 minutes.") from None

    if not visual or settings.demo_mode:
        if mime == "application/pdf" and not text.strip():
            raise HTTPException(422, "Configurez MISTRAL_API_KEY et désactivez DEMO_MODE pour analyser ce PDF sans texte.")
        extraction = (DocumentExtraction(model="local", pages=[ExtractedPage(index=0, markdown=text, method="local")],
                     warnings=["Extraction locale limitée, sans analyse des illustrations."])
                      if text else None)
        return mime, text, extraction.model_dump() if extraction else None
    try:
        if not settings.enabled(settings.mistral_api_key):
            raise HTTPException(422, "Configurez MISTRAL_API_KEY pour analyser les PDF et images avec OCR 4.1.")
        page_count = await asyncio.to_thread(lambda: len(PdfReader(io.BytesIO(content)).pages)) if mime == "application/pdf" else 1
        async with asyncio.timeout(600):
            extraction = await providers.extract_ocr(settings, content, mime, page_count=page_count)
            for page in extraction.pages:
                readable = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", page.markdown)
                if not any(character.isalnum() for character in readable):
                    if mime == "application/pdf":
                        image = await asyncio.to_thread(render_pdf_page, content, page.index)
                        image_mime = "image/png"
                    else:
                        image = await asyncio.to_thread(normalize_vision_image, content)
                        image_mime = "image/jpeg"
                    page.vision_description = await providers.describe_image(settings, image, image_mime)
                    page.method = "vision"
            if len(extraction.model_dump_json().encode("utf-8")) > 4 * 1024 * 1024:
                raise providers.ProviderError("L'extraction dépasse la limite de stockage de 4 Mo.")
            return mime, extraction_text(extraction), extraction.model_dump()
    except TimeoutError:
        raise providers.ProviderError("L'analyse du document dépasse le délai global de 10 minutes.") from None
